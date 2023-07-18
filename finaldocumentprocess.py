import os
from docx import Document
from openpyxl import Workbook


path_to_docs = './'

wb = Workbook()

ws = wb.active

#creates and writes  the headers into the sheet for the proper naming scheme 
ws.append (['Invoice ID', 'Total number of products purchased', 'Subtotal', 'Tax', 'Total'])

# initializes the grand totals for the bottom of the sheet after everything has been calculated.
grand_total_products, grand_subtotal, grand_tax, grand_total = 0, 0.0, 0.0, 0.0

for file in os.listdir(path_to_docs):

    #goes through only files that end in .docx
    if file.endswith('.docx'):
        doc=Document(os.path.join(path_to_docs, file))

        #initializes the id, total products, sub, tax and total amount
        invoice_id, total_products, subtotal, tax, total = '', 0, '', '', ''
        for paragraph in doc.paragraphs:

            #grabs invoice number/id
            if paragraph.text.startswith ('INV'):
                invoice_id = paragraph.text.replace('INV', '')

            #calculates the total amount of whichever product is purchased
            elif 'PRODUCTS' in paragraph.text:
                    total_products= sum(int(line.split(':')[1]) for line in paragraph.text.split('\n')[1:] if line)

            #this grabs the subtotal, tax and total. Changed the sub, tax and total from a string to a float point value to allow adding them together in the grand totals at the bottom of the sheet
            elif 'SUBTOTAL' in paragraph.text:
                 lines = paragraph.text.split('\n')
                 subtotal = float(lines [0].split(':')[1])
                 tax = float(lines[1].split(':')[1])
                 total = float(lines[2].split(':')[1])

        #cant believe how long it took me to figure out that the indentation on this damn line of code was 1 too far in
        #and was creating 3 lines of each item. Was actually going nuts lol. anyways this appends the invoice details into a new row
        ws.append([invoice_id, total_products, subtotal, tax, total])


        #completely unnecessary but thought it would be neat to also add the grand total of everything combined at the bottom of the sheet. hope you dont mind. 
        grand_total_products += total_products
        grand_subtotal += subtotal
        grand_tax += tax
        grand_total += total

# append grand totals to the bottom of the sheet
ws.append(['Grand Total', grand_total_products, grand_subtotal, grand_tax, grand_total])

#saves the created workbook as an .xlsx file
wb.save("Invoices_total.xlsx")